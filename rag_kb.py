
from __future__ import annotations

import hashlib
import json
import os
import pickle
import re
import unicodedata
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable, List, Optional, Tuple

from rank_bm25 import BM25Okapi
from autogenbook.retrieval.sanitize import sanitize_context_text

# Optional dependencies imported lazily inside loaders:
# - pypdf
# - python-docx
# - python-pptx


SUPPORTED_EXTS = {".pdf", ".docx", ".pptx", ".md", ".txt"}


@dataclass
class Chunk:
    id: str
    source_path: str
    loc: str  # page/slide/section hint
    text: str
    rid: str = ""
    cite_key: str = ""


def _normalize_text(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", text)
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return normalized.casefold()


def _tokenize(text: str) -> List[str]:
    normalized = _normalize_text(text)
    return re.findall(r"[a-z0-9]+", normalized)


def _env_bool(name: str, default: bool) -> bool:
    value = os.environ.get(name)
    if value is None:
        return default
    value = value.strip().lower()
    if value in {"1", "true", "yes", "on"}:
        return True
    if value in {"0", "false", "no", "off"}:
        return False
    return default


def _env_str(name: str, default: str) -> str:
    value = os.environ.get(name)
    if value is None:
        return default
    value = value.strip()
    return value or default


def _chunk_text(text: str, chunk_chars: int = 1800, overlap_chars: int = 250) -> List[str]:
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(len(text), start + chunk_chars)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == len(text):
            break
        start = max(0, end - overlap_chars)
    return chunks


def _chunk_text_by_headings(
    text: str, chunk_chars: int = 1800, overlap_chars: int = 250
) -> List[str]:
    lines = text.splitlines()
    sections: List[str] = []
    current: List[str] = []
    for line in lines:
        if re.match(r"^\s{0,3}#{1,6}\s+\S", line):
            if current:
                sections.append("\n".join(current).strip())
            current = [line.strip()]
        else:
            current.append(line)
    if current:
        sections.append("\n".join(current).strip())
    chunks: List[str] = []
    for section in sections:
        chunks.extend(_chunk_text(section, chunk_chars, overlap_chars))
    return [c for c in chunks if c]


def _read_txt_md(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(
    path: Path,
    *,
    enable_ocr: bool = False,
    ocr_lang: str = "eng",
) -> Tuple[str, List[Tuple[str, str]]]:
    """
    Returns:
      full_text,
      list of (loc, page_text)
    """
    from pypdf import PdfReader  # type: ignore

    reader = PdfReader(str(path))
    convert_from_path = None
    if enable_ocr:
        try:
            from pdf2image import convert_from_path  # type: ignore
        except Exception:
            convert_from_path = None
    page_texts: List[Tuple[str, str]] = []
    full = []
    for i, page in enumerate(reader.pages, start=1):
        t = page.extract_text() or ""
        t = t.strip()
        if not t and enable_ocr and convert_from_path is not None:
            try:
                images = convert_from_path(str(path), first_page=i, last_page=i)
                if images:
                    import pytesseract  # type: ignore

                    t = (pytesseract.image_to_string(images[0], lang=ocr_lang) or "").strip()
            except Exception:
                t = ""
        if t:
            loc = f"page {i}"
            page_texts.append((loc, t))
            full.append(t)
    return "\n".join(full), page_texts


def _read_docx(path: Path) -> str:
    from docx import Document  # type: ignore

    doc = Document(str(path))
    paras = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    return "\n".join(paras)


def _read_pptx(path: Path) -> List[Tuple[str, str]]:
    from pptx import Presentation  # type: ignore

    prs = Presentation(str(path))
    slides: List[Tuple[str, str]] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                t = (shape.text or "").strip()
                if t:
                    texts.append(t)
        slide_text = "\n".join(texts).strip()
        if slide_text:
            slides.append((f"slide {i}", slide_text))
    return slides


def _file_sha256(path: Path, block_size: int = 1 << 20) -> str:
    """
    Content hash of a file's bytes, independent of its path/filename/mtime.
    Used to key the extraction cache below so the *same document* uploaded
    under different names or into different projects/runs is only ever
    parsed once, while any change to its actual bytes is always detected.
    """
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(block_size), b""):
            h.update(block)
    return h.hexdigest()


# Bump when the on-disk payload shape or the semantics of a cached extraction
# change, so old entries already sitting in a long-lived, cross-deployment cache
# volume are transparently invalidated instead of silently going on serving
# pre-fix output forever.
_EXTRACTION_CACHE_VERSION = 1


def _extraction_cache_path(cache_dir: Path, file_sha256: str, params_tag: str) -> Path:
    key = hashlib.sha256(
        f"v{_EXTRACTION_CACHE_VERSION}:{file_sha256}:{params_tag}".encode("utf-8")
    ).hexdigest()[:40]
    return cache_dir / f"extract_{key}.json"


def _load_extraction_cache(cache_path: Path, required_key: str) -> Optional[dict]:
    """
    Loads a cached extraction payload, but only if it parses *and* has the shape
    the caller expects. A cache is only ever allowed to make extraction faster,
    never less correct than having no cache at all — so any corruption, permission
    error, or schema mismatch (an older/newer payload format, e.g. after a future
    change to this function) is treated as an ordinary cache miss rather than
    raising and silently dropping the whole file from the knowledge base.
    """
    if not cache_path.exists():
        return None
    try:
        payload = json.loads(cache_path.read_text(encoding="utf-8"))
    except Exception as e:
        print(f"[KB] Varování: poškozená cache položka {cache_path} ({e}). Ignoruji cache.")
        return None
    if not isinstance(payload, dict) or required_key not in payload:
        print(f"[KB] Varování: neočekávaný formát cache položky {cache_path}. Ignoruji cache.")
        return None
    if required_key == "text":
        return payload if isinstance(payload["text"], str) else None
    pages = payload.get("pages")
    if not isinstance(pages, list) or not all(
        isinstance(p, dict) and isinstance(p.get("loc"), str) and isinstance(p.get("text"), str)
        for p in pages
    ):
        return None
    return payload


def _store_extraction_cache(cache_path: Path, payload: dict) -> None:
    # Write-then-rename so a run reading this entry concurrently with another run
    # writing it never sees a partial file (os.replace is atomic on the same
    # filesystem, unlike writing cache_path directly). The temp suffix mixes a PID
    # with a UUID: the PID alone can collide across independently PID-namespaced
    # worker containers sharing this same cache volume.
    tmp_path = cache_path.with_name(cache_path.name + f".tmp-{os.getpid()}-{uuid.uuid4().hex}")
    try:
        tmp_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        os.replace(tmp_path, cache_path)
    except Exception:
        try:
            tmp_path.unlink(missing_ok=True)
        except Exception:
            pass


def _cached_text_extraction(
    file_cache_path: Optional[Path], force_rebuild: bool, extractor: Callable[[], str]
) -> str:
    """Shared cache-or-extract path for extractors returning one text blob (.txt/.md, .docx)."""
    if file_cache_path is not None and not force_rebuild:
        cached = _load_extraction_cache(file_cache_path, "text")
        if cached is not None:
            return cached["text"]
    text = extractor()
    if file_cache_path is not None:
        _store_extraction_cache(file_cache_path, {"text": text})
    return text


def _cached_pages_extraction(
    file_cache_path: Optional[Path],
    force_rebuild: bool,
    extractor: Callable[[], List[Tuple[str, str]]],
) -> List[Tuple[str, str]]:
    """Shared cache-or-extract path for extractors returning (loc, text) pairs (.pdf pages, .pptx slides)."""
    if file_cache_path is not None and not force_rebuild:
        cached = _load_extraction_cache(file_cache_path, "pages")
        if cached is not None:
            return [(p["loc"], p["text"]) for p in cached["pages"]]
    pages = extractor()
    if file_cache_path is not None:
        _store_extraction_cache(
            file_cache_path, {"pages": [{"loc": loc, "text": text} for loc, text in pages]}
        )
    return pages


def _dir_fingerprint(dir_path: Path) -> str:
    """
    Hash of relative paths + sizes + mtimes for supported files.
    """
    h = hashlib.sha256()
    for p in sorted(dir_path.rglob("*")):
        if not p.is_file():
            continue
        if p.suffix.lower() not in SUPPORTED_EXTS:
            continue
        stat = p.stat()
        rel = p.relative_to(dir_path).as_posix()
        h.update(rel.encode("utf-8", errors="ignore"))
        h.update(str(stat.st_size).encode())
        h.update(str(int(stat.st_mtime_ns)).encode())
    return h.hexdigest()


def _clean_key(text: str) -> str:
    return re.sub(r"[^a-zA-Z0-9]+", "_", text.strip().lower()).strip("_")


def _sanitize_key(text: str) -> str:
    cleaned = _clean_key(text)
    return cleaned or "unknown"


def _sanitize_loc(loc: str) -> str:
    cleaned = _clean_key(loc)
    return cleaned or "loc"


def _source_id_from_path(path: Path, root: Optional[Path] = None, max_len: int = 48) -> str:
    rel = path
    if root is not None:
        try:
            rel = path.relative_to(root)
        except ValueError:
            rel = path
    rel_str = rel.as_posix()
    base = _clean_key(rel_str)
    if max_len > 0 and len(base) > max_len:
        base = base[:max_len].rstrip("_")
    digest = hashlib.sha256(rel_str.encode("utf-8", errors="ignore")).hexdigest()[:8]
    if not base:
        base = "source"
    return f"{base}_{digest}"


def _make_rid(source_id: str, loc: str, chunk_index: int) -> str:
    loc_key = _sanitize_loc(loc)
    return f"RID:kb:{source_id}:{loc_key}:{chunk_index}"


def _make_cite_key(source_id: str, loc: str, chunk_index: int) -> str:
    return f"kb_{source_id}_{_sanitize_key(loc)}_{chunk_index}"


class KnowledgeBase:
    """
    Lightweight RAG knowledge base using BM25 over chunked documents.
    Designed to be robust, local-only (no embeddings), and fast enough for CLI.

    - Build from a directory containing PDF/DOCX/PPTX/MD/TXT.
    - Retrieve top-k chunks for a query string.
    - Cache to disk to avoid rebuilding.
    """

    def __init__(self, chunks: List[Chunk], bm25: BM25Okapi):
        self.chunks = chunks
        self._bm25 = bm25
        self._tokenized_corpus = None  # just for debugging if needed
        self._ensure_chunk_ids()

    @staticmethod
    def build_from_directory(
        dir_path: Path,
        cache_dir: Optional[Path] = None,
        force_rebuild: bool = False,
        chunk_chars: int = 1800,
        overlap_chars: int = 250,
        enable_ocr: Optional[bool] = None,
        heading_chunks: Optional[bool] = None,
        ocr_lang: Optional[str] = None,
        extract_cache_dir: Optional[Path] = None,
    ) -> "KnowledgeBase":
        dir_path = dir_path.expanduser().resolve()
        if not dir_path.exists() or not dir_path.is_dir():
            raise FileNotFoundError(f"KB adresář neexistuje nebo není adresář: {dir_path}")

        cache_dir = cache_dir or (Path.cwd() / ".autogenbook_kb_cache")
        cache_dir.mkdir(parents=True, exist_ok=True)

        # Resolved before the whole-directory cache lookup below so a change to any of
        # them (independent of the directory's own contents) correctly invalidates that
        # cache instead of it silently returning a `KnowledgeBase` built under the old
        # settings.
        enable_ocr = enable_ocr if enable_ocr is not None else _env_bool("AUTOGENBOOK_KB_OCR", False)
        heading_chunks = (
            heading_chunks if heading_chunks is not None else _env_bool("AUTOGENBOOK_KB_HEADING_CHUNKS", False)
        )
        ocr_lang = ocr_lang or _env_str("AUTOGENBOOK_KB_OCR_LANG", "eng")

        fp = _dir_fingerprint(dir_path)
        config_tag = f"ocr={int(enable_ocr)}:lang={ocr_lang}:heading={int(heading_chunks)}"
        combined_fp = hashlib.sha256(f"{fp}:{config_tag}".encode("utf-8")).hexdigest()
        cache_key = hashlib.sha256(str(dir_path).encode("utf-8")).hexdigest()[:16]
        meta_path = cache_dir / f"kb_{cache_key}.meta.json"
        data_path = cache_dir / f"kb_{cache_key}.pkl"

        if not force_rebuild and meta_path.exists() and data_path.exists():
            try:
                meta = json.loads(meta_path.read_text(encoding="utf-8"))
                if meta.get("fingerprint") == combined_fp:
                    with open(data_path, "rb") as f:
                        obj = pickle.load(f)
                    if isinstance(obj, KnowledgeBase):
                        obj._ensure_chunk_ids()
                    return obj
            except Exception:
                # fall through to rebuild
                pass

        # Per-file extraction cache: keyed by file *content* hash rather than path, so the
        # same document re-uploaded under a different name, or attached to a different
        # project/run, is only ever parsed/OCR'd once. Disabled unless a directory is
        # configured (a plain `--kb-dir` CLI run has none set and behaves exactly as
        # before); the web app points this at a persistent, cross-run volume distinct
        # from the per-run `cache_dir` above, which lives inside an ephemeral work dir.
        if extract_cache_dir is None:
            configured = _env_str("AUTOGENBOOK_KB_EXTRACT_CACHE_DIR", "")
            extract_cache_dir = Path(configured).expanduser().resolve() if configured else None
        if extract_cache_dir is not None:
            try:
                extract_cache_dir.mkdir(parents=True, exist_ok=True)
            except OSError as e:
                print(
                    f"[KB] Varování: nelze vytvořit extract cache adresář {extract_cache_dir} "
                    f"({e}). Extract cache je pro tento běh vypnutá."
                )
                extract_cache_dir = None

        chunks: List[Chunk] = []
        for file_path in sorted(dir_path.rglob("*")):
            if not file_path.is_file():
                continue
            ext = file_path.suffix.lower()
            if ext not in SUPPORTED_EXTS:
                continue

            try:
                source_id = _source_id_from_path(file_path, dir_path)
                file_cache_path = None
                # PDF+OCR is excluded: OCR is an external, non-deterministic process
                # (rendering pages to images, then Tesseract), and a transient failure
                # there must not be baked into this persistent, cross-run/cross-project
                # cache forever — unlike the other extractors, which are pure functions
                # of the file's bytes and always safe to cache.
                cache_eligible = extract_cache_dir is not None and not (ext == ".pdf" and enable_ocr)
                if cache_eligible:
                    file_sha256 = _file_sha256(file_path)
                    params_tag = f"{ext}:ocr={int(enable_ocr)}:lang={ocr_lang}" if ext == ".pdf" else ext
                    file_cache_path = _extraction_cache_path(extract_cache_dir, file_sha256, params_tag)

                if ext in {".txt", ".md"}:
                    text = _cached_text_extraction(
                        file_cache_path, force_rebuild, lambda: _read_txt_md(file_path)
                    )
                    if heading_chunks and ext == ".md":
                        parts = _chunk_text_by_headings(text, chunk_chars, overlap_chars)
                    else:
                        parts = _chunk_text(text, chunk_chars, overlap_chars)
                    for j, chunk in enumerate(parts, start=1):
                        loc_base = "chunk"
                        rid = _make_rid(source_id, loc_base, j)
                        cite_key = _make_cite_key(source_id, loc_base, j)
                        chunks.append(
                            Chunk(
                                id=f"kb:{source_id}:{loc_base}:{j}",
                                source_path=str(file_path),
                                loc=f"chunk {j}",
                                text=chunk,
                                rid=rid,
                                cite_key=cite_key,
                            )
                        )

                elif ext == ".pdf":
                    pages = _cached_pages_extraction(
                        file_cache_path,
                        force_rebuild,
                        lambda: _read_pdf(file_path, enable_ocr=enable_ocr, ocr_lang=ocr_lang)[1],
                    )
                    for loc, page_text in pages:
                        for j, chunk in enumerate(_chunk_text(page_text, chunk_chars, overlap_chars), start=1):
                            rid = _make_rid(source_id, loc, j)
                            cite_key = _make_cite_key(source_id, loc, j)
                            chunks.append(
                                Chunk(
                                    id=f"kb:{source_id}:{loc}:{j}",
                                    source_path=str(file_path),
                                    loc=f"{loc}, chunk {j}",
                                    text=chunk,
                                    rid=rid,
                                    cite_key=cite_key,
                                )
                            )

                elif ext == ".docx":
                    text = _cached_text_extraction(
                        file_cache_path, force_rebuild, lambda: _read_docx(file_path)
                    )
                    for j, chunk in enumerate(_chunk_text(text, chunk_chars, overlap_chars), start=1):
                        loc_base = "chunk"
                        rid = _make_rid(source_id, loc_base, j)
                        cite_key = _make_cite_key(source_id, loc_base, j)
                        chunks.append(
                            Chunk(
                                id=f"kb:{source_id}:{loc_base}:{j}",
                                source_path=str(file_path),
                                loc=f"chunk {j}",
                                text=chunk,
                                rid=rid,
                                cite_key=cite_key,
                            )
                        )

                elif ext == ".pptx":
                    slides = _cached_pages_extraction(
                        file_cache_path, force_rebuild, lambda: _read_pptx(file_path)
                    )
                    for loc, slide_text in slides:
                        for j, chunk in enumerate(_chunk_text(slide_text, chunk_chars, overlap_chars), start=1):
                            rid = _make_rid(source_id, loc, j)
                            cite_key = _make_cite_key(source_id, loc, j)
                            chunks.append(
                                Chunk(
                                    id=f"kb:{source_id}:{loc}:{j}",
                                    source_path=str(file_path),
                                    loc=f"{loc}, chunk {j}",
                                    text=chunk,
                                    rid=rid,
                                    cite_key=cite_key,
                                )
                            )
            except Exception as e:
                # Continue on errors (e.g., scanned PDFs with no text).
                # You can inspect and add OCR later if needed.
                print(f"[KB] Varování: nepodařilo se načíst {file_path} ({e}). Přeskakuji.")

        tokenized_corpus = [_tokenize(c.text) for c in chunks]
        bm25 = BM25Okapi(tokenized_corpus)

        kb = KnowledgeBase(chunks=chunks, bm25=bm25)

        # Cache
        try:
            with open(data_path, "wb") as f:
                pickle.dump(kb, f)
            meta_path.write_text(
                json.dumps({"dir": str(dir_path), "fingerprint": combined_fp}, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
        except Exception:
            pass

        return kb

    def _ensure_chunk_ids(self) -> None:
        for chunk in self.chunks:
            source_id = _source_id_from_path(Path(chunk.source_path))
            loc = getattr(chunk, "loc", "chunk")
            loc_base = loc.split(",")[0].strip() if isinstance(loc, str) and loc else "chunk"
            if loc_base.lower().startswith("chunk"):
                loc_base = "chunk"
            idx = _extract_chunk_index(loc) or _extract_chunk_index(chunk.id) or 1
            if not getattr(chunk, "rid", ""):
                chunk.rid = _make_rid(source_id, loc_base, idx)
            if not getattr(chunk, "cite_key", ""):
                chunk.cite_key = _make_cite_key(source_id, loc_base, idx)

    def retrieve(self, query: str, k: int = 6) -> List[Tuple[Chunk, float]]:
        if not self.chunks:
            return []
        q_tokens = _tokenize(query)
        scores = self._bm25.get_scores(q_tokens)
        # top-k indices
        top_idx = sorted(range(len(scores)), key=lambda i: scores[i], reverse=True)[:k]
        results: List[Tuple[Chunk, float]] = []
        for i in top_idx:
            if scores[i] <= 0:
                continue
            results.append((self.chunks[i], float(scores[i])))
        return results

    def format_context(self, query: str, k: int = 6, max_chars_total: int = 6000) -> str:
        """
        Returns a compact, prompt-ready context block with top-k chunks.
        """
        results = self.retrieve(query, k=k)
        if not results:
            return ""

        blocks = []
        remaining = max_chars_total
        for _, (chunk, score) in enumerate(results, start=1):
            # keep excerpt bounded
            excerpt = sanitize_context_text(chunk.text.strip())
            excerpt = excerpt[: min(len(excerpt), 1500)]
            if not chunk.rid or not chunk.cite_key:
                self._ensure_chunk_ids()
            rid = chunk.rid
            cite_key = chunk.cite_key
            header = (
                f"[{rid}] kind=kb source=\"{Path(chunk.source_path).name}\" "
                f"loc=\"{chunk.loc}\" score={score:.2f} cite_key=\"{cite_key}\""
            )
            block = f"{header}\n{excerpt}"
            if len(block) + 2 > remaining:
                break
            blocks.append(block)
            remaining -= len(block) + 2

        return "\n\n".join(blocks)


def _extract_chunk_index(text: str) -> Optional[int]:
    if not text:
        return None
    match = re.search(r"chunk\s+(\d+)", text, flags=re.IGNORECASE)
    if match:
        return int(match.group(1))
    match = re.search(r":(\d+)$", text)
    if match:
        return int(match.group(1))
    return None
