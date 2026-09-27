"""Slide deck outputs: the deck Markdown (the old engine's format), PPTX and
Beamer LaTeX.

Deck Markdown, as `autogenbook/pipelines/presentation_pipeline.py:
_build_presentation_markdown` wrote it:

    theme: Madrid
    paginate: false
    outline: true
    author: ...            (header:/footer: when set)

    ---
    <!-- _class: title -->
    # Title
    <summary, author>

    ---
    # Slide title
    <slide body>

PPTX ports `autogenbook/presentation_export.py:md_to_pptx` (python-pptx,
16:9, title + content layouts) and adds the narration as speaker notes.
Beamer is produced by pandoc (`-t beamer`, one frame per slide, a section
per slide for the outline frame) and compiled with LuaLaTeX instead of the
old hand-built preamble compiled with pdfLaTeX.
"""

from __future__ import annotations

import re
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping, Sequence

from engine.assemble.latex import LatexError, babel_language, pandoc_available

TITLE_MARKER = "<!-- _class: title -->"
MAX_TEXT_LINES = 10
_FRAME_SPLIT = "\n---\n"
_HEADING_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$")
_RULE_RE = re.compile(r"^\s{0,3}([-*_])(\s*\1){2,}\s*$")
_IMAGE_RE = re.compile(r"!\[\[(.*?)\]\]|!\[[^\]]*\]\([^)]*\)")
_LATEX_SPECIAL = {"\\": r"\textbackslash{}", "{": r"\{", "}": r"\}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}


@dataclass
class Frame:
    index: int  # 1-based position in the deck; the title slide is 1
    title: str
    body: str
    is_title: bool = False


OUTLINE_TITLES = {"cs": "Osnova", "sk": "Osnova", "de": "Gliederung", "pl": "Plan", "fr": "Plan", "es": "Índice"}


@dataclass
class DeckSettings:
    title: str
    summary: str = ""
    author: str = ""
    theme: str = "Madrid"
    paginate: bool = False
    outline: bool = True
    header: str = ""
    footer: str = ""
    language: str = "en"

    @classmethod
    def from_attrs(cls, title: str, summary: str, author: str, attrs: Mapping[str, Any]) -> "DeckSettings":
        theme = str(attrs.get("theme") or "Madrid").strip()
        return cls(
            title=title, summary=summary, author=author,
            theme="Madrid" if not theme or theme.lower() == "beamer" else theme,
            paginate=bool(attrs.get("paginate", False)), outline=bool(attrs.get("outline", True)),
            header=str(attrs.get("header") or "").strip(), footer=str(attrs.get("footer") or "").strip(),
        )


def slide_body(text: str) -> str:
    """A slide body safe inside the deck: no headings (they would start a new
    frame), no horizontal rules (the frame separator), no image embeds."""
    lines: list[str] = []
    for raw in (text or "").strip().splitlines():
        if _RULE_RE.match(raw):
            continue
        heading = _HEADING_RE.match(raw)
        if heading:
            if lines and lines[-1].strip():
                lines.append("")
            lines += [f"**{heading.group(2).strip()}**", ""]
            continue
        line = _IMAGE_RE.sub("", raw).rstrip()
        if line.strip() or (lines and lines[-1].strip()):
            lines.append(line)
    return "\n".join(lines).strip()


def limit_lines(body: str, max_lines: int = MAX_TEXT_LINES) -> tuple[str, int]:
    """At most `max_lines` non-empty lines (the old engine's slide limit);
    returns the body and the number of lines dropped."""
    kept: list[str] = []
    count = dropped = 0
    for line in body.splitlines():
        if line.strip():
            count += 1
            if count > max_lines:
                dropped += 1
                continue
        kept.append(line)
    return "\n".join(kept).strip(), dropped


def build_deck(settings: DeckSettings, slides: Sequence[tuple[str, str]]) -> str:
    lines = [f"theme: {settings.theme}", f"paginate: {'true' if settings.paginate else 'false'}",
             f"outline: {'true' if settings.outline else 'false'}"]
    if settings.author:
        lines.append(f"author: {settings.author}")
    if settings.header:
        lines.append(f"header: {settings.header}")
    if settings.footer:
        lines.append(f"footer: {settings.footer}")
    lines += ["", "---", TITLE_MARKER, f"# {settings.title}"]
    if settings.summary:
        lines += ["", settings.summary.strip()]
    if settings.author:
        lines += ["", f"**{settings.author}**"]
    for title, body in slides:
        lines += ["", "---", f"# {title.strip()}"]
        body = slide_body(body)
        if body:
            lines += ["", body]
    return "\n".join(lines).rstrip() + "\n"


def split_frames(deck: str) -> list[Frame]:
    """The deck's frames (front matter excluded), numbered from 1."""
    frames: list[Frame] = []
    for index, raw in enumerate(deck.replace("\r\n", "\n").split(_FRAME_SPLIT)[1:], start=1):
        is_title = TITLE_MARKER in raw
        title = ""
        body_lines: list[str] = []
        for line in raw.replace(TITLE_MARKER, "").splitlines():
            match = re.match(r"^\s*#\s+(.+)$", line)
            if match and not title:
                title = match.group(1).strip()
                continue
            body_lines.append(line)
        frames.append(Frame(index, title, "\n".join(body_lines).strip(), is_title))
    return frames


def front_matter(deck: str) -> dict[str, str]:
    head = deck.split(_FRAME_SPLIT, 1)[0]
    out: dict[str, str] = {}
    for line in head.splitlines():
        match = re.match(r"^(\w+)\s*:\s*(.+)$", line.strip())
        if match:
            out[match.group(1).lower()] = match.group(2).strip()
    return out


# -------------------------------------------------------------------- PPTX
def _plain(text: str) -> str:
    cleaned = _IMAGE_RE.sub("", text.strip())
    cleaned = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1 (\2)", cleaned)
    cleaned = re.sub(r"`{1,3}(.+?)`{1,3}", r"\1", cleaned)
    cleaned = re.sub(r"(\*\*|__)(.*?)\1", r"\2", cleaned)
    cleaned = re.sub(r"(?<![*\w])\*(?!\*)(.+?)\*(?!\w)", r"\1", cleaned)
    cleaned = re.sub(r"(?<![_\w])_(?!_)(.+?)_(?!\w)", r"\1", cleaned)
    cleaned = re.sub(r"\$([^$]+)\$", r"\1", cleaned)
    return re.sub(r"\s+", " ", cleaned).strip()


def pptx_lines(body: str) -> list[tuple[str, int]]:
    """(text, indent level) per non-empty line, bullets as `•`."""
    out: list[tuple[str, int]] = []
    for raw in body.splitlines():
        if not raw.strip() or re.match(r"^\s*\|?\s*:?-{3,}", raw):
            continue
        indent = len(raw) - len(raw.lstrip(" "))
        level = min(indent // 2, 4)
        stripped = raw.strip()
        numbered = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
        bulleted = re.match(r"^[-*+]\s+(.*)$", stripped)
        if numbered:
            text = f"{numbered.group(1)}. {_plain(numbered.group(2))}"
        elif bulleted:
            text = f"• {_plain(bulleted.group(1))}"
        elif stripped.startswith("|"):
            text = " | ".join(_plain(c) for c in stripped.strip("|").split("|"))
        else:
            text = _plain(stripped)
        if text:
            out.append((text, level))
    return out


def deck_to_pptx(deck: str, out_path: Path, *, notes: Mapping[int, str] | None = None) -> Path:
    try:
        from pptx import Presentation  # type: ignore[import-not-found]
        from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN  # type: ignore[import-not-found]
        from pptx.util import Inches, Pt  # type: ignore[import-not-found]
    except ImportError as exc:
        raise RuntimeError("PPTX export needs python-pptx (pip install python-pptx)") from exc
    frames = split_frames(deck)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    title_layout, content_layout = prs.slide_layouts[0], prs.slide_layouts[1]

    def fill(text_frame: Any, lines: list[tuple[str, int]], size: int) -> None:
        text_frame.clear()
        text_frame.word_wrap = True
        text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for position, (text, level) in enumerate(lines or [("", 0)]):
            paragraph = text_frame.paragraphs[0] if position == 0 else text_frame.add_paragraph()
            paragraph.text = text
            paragraph.level = level
            paragraph.font.size = Pt(size)
            paragraph.alignment = PP_ALIGN.LEFT

    for frame in frames:
        lines = pptx_lines(frame.body)
        if frame.is_title:
            slide = prs.slides.add_slide(title_layout)
            slide.shapes.title.text = _plain(frame.title)
            slide.placeholders[1].text = "\n".join(text for text, _level in lines)
        else:
            slide = prs.slides.add_slide(content_layout)
            slide.shapes.title.text = _plain(frame.title)
            fill(slide.placeholders[1].text_frame, lines, 18 if len(lines) <= 8 else 14)
        note = (notes or {}).get(frame.index)
        if note:
            slide.notes_slide.notes_text_frame.text = note
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


# ------------------------------------------------------------------ Beamer
def settings_outline_title(settings: DeckSettings) -> str:
    return OUTLINE_TITLES.get(settings.language, "Outline")


def latex_escape(text: str) -> str:
    return "".join(_LATEX_SPECIAL.get(ch, ch) for ch in text)


def beamer_markdown(deck: str, settings: DeckSettings) -> str:
    """Pandoc input: one `#` section + `##` frame per slide (the outline frame
    lists the slides, as the old `\\section{}` per frame did)."""
    parts: list[str] = []
    if settings.outline:
        outline = settings_outline_title(settings)
        parts.append(f"## {outline} {{.allowframebreaks}}\n\n```{{=latex}}\n\\tableofcontents\n```\n")
    for frame in split_frames(deck):
        if frame.is_title:
            continue
        title = frame.title.replace("{", "(").replace("}", ")")
        attrs = " {.allowframebreaks}" if (frame.body.count("\n") >= MAX_TEXT_LINES or "|" in frame.body) else ""
        parts.append(f"# {title}\n\n## {title}{attrs}\n\n{frame.body}\n")
    return "\n".join(parts)


def _footline(settings: DeckSettings) -> str:
    if not (settings.header or settings.footer):
        return r"\setbeamertemplate{footline}[frame number]" if settings.paginate else ""
    pages = r"\insertframenumber{} / \inserttotalframenumber" if settings.paginate else ""
    return (
        "\\setbeamertemplate{footline}{\\leavevmode\\hbox{%\n"
        "\\begin{beamercolorbox}[wd=.45\\paperwidth,ht=2.25ex,dp=1ex,leftskip=1em]{author in head/foot}"
        f"{latex_escape(settings.header)}\\end{{beamercolorbox}}%\n"
        "\\begin{beamercolorbox}[wd=.45\\paperwidth,ht=2.25ex,dp=1ex,center]{title in head/foot}"
        f"{latex_escape(settings.footer)}\\end{{beamercolorbox}}%\n"
        "\\begin{beamercolorbox}[wd=.1\\paperwidth,ht=2.25ex,dp=1ex,right,rightskip=1em]{date in head/foot}"
        f"{pages}\\end{{beamercolorbox}}}}\\vskip0pt}}"
    )


def deck_to_beamer(deck: str, out_tex: Path, *, settings: DeckSettings, language: str) -> Path:
    if not pandoc_available():
        raise LatexError("pandoc is required for Beamer export (install pandoc)")
    extensions = "markdown+tex_math_dollars+pipe_tables-implicit_figures-auto_identifiers-raw_tex"
    args = [
        "pandoc", f"--from={extensions}", "--to=beamer", "--standalone", "--slide-level=2", "--no-highlight", "--wrap=none",
        "-M", f"title={settings.title}", "-M", "date=", "-V", f"theme={settings.theme}", "-V", "aspectratio=169",
        "-M", "section-titles=false", "-V", f"lang={language}", "-V", f"babel-lang={babel_language(language)}",
    ]
    if settings.author:
        args += ["-M", f"author={settings.author}"]
    args += ["-V", r"header-includes=\setbeamertemplate{frametitle continuation}{}"]
    footline = _footline(settings)
    if footline:
        args += ["-V", f"header-includes={footline}"]
    proc = subprocess.run(
        args, input=beamer_markdown(deck, settings), capture_output=True, text=True, encoding="utf-8", errors="replace",
        timeout=600, check=False,
    )
    if proc.returncode != 0:
        raise LatexError(f"pandoc failed: {(proc.stderr or proc.stdout).strip()[:1000]}")
    out_tex.parent.mkdir(parents=True, exist_ok=True)
    out_tex.write_text(proc.stdout, encoding="utf-8")
    return out_tex
