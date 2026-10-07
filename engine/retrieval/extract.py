"""Structure-aware text extraction with the shared content-hash cache.

- PDF: pypdfium2 text rectangles grouped into lines and paragraphs by layout
  (vertical gaps, font size), headings detected by font size, line-break
  hyphenation repaired (including pdfium's U+FFFE soft hyphen), running
  headers/footers dropped (lines repeated on >= 3 pages at the page top or
  bottom), tables found with pdfplumber and kept whole as Markdown tables.
  Optional OCR for text-less pages (AUTOGENBOOK_KB_OCR). PyMuPDF is not used
  (AGPL).
- DOCX (python-docx, heading styles, tables), PPTX (python-pptx, one section
  per slide), Markdown (headings, fenced code, tables, lists) and TXT.

pdfium is not thread-safe, so every pdfium call holds the extractor's lock;
the rest of the per-file work (pdfplumber, DOCX/PPTX/MD parsing) runs in
parallel in the KB builder's thread pool.

Cache: `AUTOGENBOOK_KB_EXTRACT_CACHE_DIR` (the directory the old engine
already uses, shared across runs and projects), keyed by file content hash +
extraction parameters. The key scheme carries its own `engine-v1` tag, so the
old engine's entries are plain misses and never read; writes are atomic and a
corrupt entry is a miss.
"""

from __future__ import annotations

import hashlib
import json
import re
import statistics
import threading
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from engine.retrieval.types import Block, ExtractedDoc
from engine.util.fs import atomic_write_text, sha256_file

SUPPORTED_EXTS = {".pdf", ".docx", ".pptx", ".md", ".txt"}
EXTRACTOR_VERSION = "engine-v1"

_SOFT_HYPHEN = "\ufffe"  # pdfium marks line-break hyphenation with U+FFFE / \x02


@dataclass
class ExtractOptions:
    ocr: bool = False
    ocr_lang: str = "eng"
    cache_dir: Path | None = None
    force: bool = False

    def params_tag(self, ext: str) -> str:
        if ext == ".pdf":
            return f"{EXTRACTOR_VERSION}:{ext}:ocr={int(self.ocr)}:lang={self.ocr_lang}"
        return f"{EXTRACTOR_VERSION}:{ext}"


@dataclass
class Extractor:
    options: ExtractOptions
    pdfium_lock: threading.Lock = field(default_factory=threading.Lock)
    cache_hits: int = 0
    cache_misses: int = 0

    # ----------------------------------------------------------------- cache
    def cache_path(self, file_sha256: str, ext: str) -> Path | None:
        if self.options.cache_dir is None:
            return None
        key = hashlib.sha256(f"{file_sha256}:{self.options.params_tag(ext)}".encode("utf-8")).hexdigest()[:40]
        return self.options.cache_dir / f"extract_{key}.json"

    def _load_cache(self, path: Path | None) -> list[Block] | None:
        if path is None or self.options.force or not path.exists():
            return None
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
            if payload.get("extractor") != EXTRACTOR_VERSION or not isinstance(payload.get("blocks"), list):
                return None
            return [Block.from_json(b) for b in payload["blocks"]]
        except (OSError, ValueError, KeyError, TypeError):
            return None

    def _store_cache(self, path: Path | None, blocks: list[Block]) -> None:
        if path is None:
            return
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            atomic_write_text(path, json.dumps({"extractor": EXTRACTOR_VERSION, "blocks": [b.to_json() for b in blocks]}, ensure_ascii=False))
        except OSError:
            pass  # a cache may only ever make extraction faster, never fail it

    # ------------------------------------------------------------------ main
    def extract(self, path: Path, root: Path) -> ExtractedDoc:
        ext = path.suffix.lower()
        rel = path.relative_to(root).as_posix() if path.is_relative_to(root) else path.name
        doc = ExtractedDoc(path=str(path), rel_path=rel, blocks=[])
        # PDF + OCR is not cached: OCR is external and non-deterministic, and a
        # transient failure must not be baked into a cross-run cache.
        cacheable = not (ext == ".pdf" and self.options.ocr)
        cache_path = self.cache_path(sha256_file(path), ext) if cacheable else None
        cached = self._load_cache(cache_path)
        if cached is not None:
            self.cache_hits += 1
            doc.blocks = cached
            return doc
        self.cache_misses += 1
        if ext == ".pdf":
            doc.blocks = self._pdf(path, doc.warnings)
        elif ext == ".docx":
            doc.blocks = _docx(path)
        elif ext == ".pptx":
            doc.blocks = _pptx(path)
        elif ext == ".md":
            doc.blocks = parse_markdown(path.read_text(encoding="utf-8", errors="replace"))
        elif ext == ".txt":
            doc.blocks = parse_plain_text(path.read_text(encoding="utf-8", errors="replace"))
        if cacheable:
            self._store_cache(cache_path, doc.blocks)
        return doc

    # ------------------------------------------------------------------- PDF
    def _pdf(self, path: Path, warnings: list[str]) -> list[Block]:
        import pypdfium2 as pdfium
        import pypdfium2.raw as pdfium_c

        pages: list[list[_Line]] = []
        table_candidates: list[int] = []
        page_sizes: list[tuple[float, float]] = []
        ocr_pages: list[int] = []
        with self.pdfium_lock:
            pdf = pdfium.PdfDocument(str(path))
            try:
                for index in range(len(pdf)):
                    page = pdf[index]
                    width, height = page.get_size()
                    page_sizes.append((width, height))
                    textpage = page.get_textpage()
                    lines = _pdf_lines(textpage, pdfium_c, height)
                    if not lines and self.options.ocr:
                        ocr_pages.append(index)
                    if _looks_like_table_page(page, pdfium_c):
                        table_candidates.append(index)
                    pages.append(lines)
                    textpage.close()
                    page.close()
                ocr_images = {i: pdf[i].render(scale=2.0).to_pil() for i in ocr_pages}
            finally:
                pdf.close()
        for index, image in ocr_images.items():
            text = _ocr(image, self.options.ocr_lang, warnings)
            if text:
                pages[index] = [_Line(text=t, top=float(n), bottom=float(n) + 1, size=10.0, left=0.0) for n, t in enumerate(text.splitlines()) if t.strip()]
        tables = _pdf_tables(path, table_candidates, warnings) if table_candidates else {}
        _drop_running_lines(pages)
        return _pdf_blocks(pages, tables, page_sizes)


# --------------------------------------------------------------------- PDF
@dataclass
class _Line:
    text: str
    top: float  # distance from the page top
    bottom: float
    size: float
    left: float
    page: int = 0

    @property
    def height(self) -> float:
        return max(0.1, self.bottom - self.top)


def _pdf_lines(textpage: Any, pdfium_c: Any, page_height: float) -> list[_Line]:
    rects: list[tuple[float, float, float, float, str, float]] = []
    for i in range(textpage.count_rects()):
        left, bottom, right, top = textpage.get_rect(i)
        text = textpage.get_text_bounded(left, bottom, right, top)
        if not text or not text.strip():
            continue
        index = textpage.get_index(left + 0.5, (top + bottom) / 2, 2, 2)
        size = float(pdfium_c.FPDFText_GetFontSize(textpage.raw, index)) if index is not None and index >= 0 else 0.0
        text = text.replace("\r", "").replace("\n", " ").replace("\x02", _SOFT_HYPHEN).replace("\u00ad", _SOFT_HYPHEN)
        rects.append((left, page_height - top, right, page_height - bottom, text, size or (top - bottom)))
    lines: list[_Line] = []
    for left, top, right, bottom, text, size in rects:
        middle = (top + bottom) / 2
        if lines:
            last = lines[-1]
            if abs(((last.top + last.bottom) / 2) - middle) < 0.5 * max(last.height, bottom - top) and left >= last.left:
                last.text = f"{last.text} {text}".strip()
                last.size = max(last.size, size)
                last.bottom = max(last.bottom, bottom)
                continue
        lines.append(_Line(text=text.strip(), top=top, bottom=bottom, size=size, left=left))
    return lines


def _looks_like_table_page(page: Any, pdfium_c: Any) -> bool:
    """Only pages with ruling-like path objects go to pdfplumber (slow)."""
    try:
        paths = sum(1 for _ in page.get_objects(filter=[pdfium_c.FPDF_PAGEOBJ_PATH], max_depth=2))
    except Exception:  # noqa: BLE001
        return False
    return paths >= 3


def _pdf_tables(path: Path, pages: list[int], warnings: list[str]) -> dict[int, list[tuple[float, float, str | None]]]:
    """page index -> [(top, bottom, markdown | None)]. Ruled grids come back as
    Markdown tables from pdfplumber; open tables (booktabs-style horizontal
    rules only) come back as a region (markdown None) whose lines are kept
    together as one table block."""
    try:
        import pdfplumber
    except ImportError:
        return {}
    out: dict[int, list[tuple[float, float, str | None]]] = {}
    try:
        with pdfplumber.open(str(path)) as pdf:
            for index in pages:
                if index >= len(pdf.pages):
                    continue
                page = pdf.pages[index]
                try:
                    regions = _page_table_regions(page)
                finally:
                    # pdfplumber caches every parsed object (chars, lines, rects, images)
                    # on the Page for the life of the document, so without this a
                    # 700-page textbook grows to ~3 GB and OOM-kills the worker (#178).
                    page.close()
                if regions:
                    out[index] = regions
    except Exception as exc:  # noqa: BLE001 - tables are an enhancement, text still counts
        warnings.append(f"table extraction failed: {exc}")
    return out


def _page_table_regions(page: Any) -> list[tuple[float, float, str | None]]:
    """One page's `(top, bottom, markdown | None)` regions: pdfplumber tables as
    Markdown, else ruled regions (markdown None) from horizontal rules."""
    regions: list[tuple[float, float, str | None]] = []
    for table in page.find_tables():
        markdown = table_to_markdown(table.extract())
        if markdown:
            _x0, top, _x1, bottom = table.bbox
            regions.append((float(top), float(bottom), markdown))
    if regions:
        return regions
    rules: dict[tuple[int, int], list[float]] = {}
    for line in list(page.lines) + [r for r in page.rects if abs(r["bottom"] - r["top"]) < 1.5]:
        if abs(line["bottom"] - line["top"]) > 1.5 or line["x1"] - line["x0"] < 100:
            continue
        rules.setdefault((round(line["x0"] / 3), round(line["x1"] / 3)), []).append(float(line["top"]))
    for tops in rules.values():
        if len(tops) >= 2 and max(tops) - min(tops) > 10:
            regions.append((min(tops), max(tops), None))
    return regions


def table_to_markdown(rows: list[list[Any]]) -> str:
    cleaned = [[re.sub(r"\s+", " ", str(cell or "")).strip() for cell in row] for row in rows if row]
    cleaned = [row for row in cleaned if any(row)]
    if len(cleaned) < 2:
        return ""
    width = max(len(row) for row in cleaned)
    cleaned = [row + [""] * (width - len(row)) for row in cleaned]
    lines = ["| " + " | ".join(cleaned[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(row) + " |" for row in cleaned[1:]]
    return "\n".join(lines)


def _running_key(text: str) -> str:
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", text)).strip().casefold()


def _drop_running_lines(pages: list[list[_Line]], window: int = 2, min_pages: int = 3, y_tolerance: float = 6.0) -> None:
    """Remove running headers/footers: lines within the first/last `window`
    lines of a page whose digit-normalised text recurs at (about) the same
    vertical position on >= `min_pages` pages, and bare page numbers."""
    positions: dict[str, list[float]] = {}
    for lines in pages:
        for line in {id(l): l for l in lines[:window] + lines[-window:]}.values():
            positions.setdefault(_running_key(line.text), []).append(line.top)
    running: dict[str, tuple[float, float]] = {}
    for key, tops in positions.items():
        if not key or len(tops) < min_pages:
            continue
        tops.sort()
        best: list[float] = []
        cluster = [tops[0]]
        for top in tops[1:]:
            if top - cluster[-1] <= y_tolerance:
                cluster.append(top)
            else:
                best = max(best, cluster, key=len)
                cluster = [top]
        best = max(best, cluster, key=len)
        if len(best) >= min_pages:
            running[key] = (best[0] - y_tolerance, best[-1] + y_tolerance)
    for p, lines in enumerate(pages):
        keep: list[_Line] = []
        for i, line in enumerate(lines):
            at_edge = i < window or i >= len(lines) - window
            key = _running_key(line.text)
            band = running.get(key)
            if at_edge and (
                (band is not None and band[0] <= line.top <= band[1])
                or re.fullmatch(r"(?:page|strana|str\.)?\s*#(?:\s*/\s*#)?", key)
            ):
                continue
            keep.append(line)
        pages[p] = keep


def dehyphenate_join(lines: list[str]) -> str:
    text = ""
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if not text:
            text = line
        elif text.rstrip().endswith(_SOFT_HYPHEN):
            text = text.rstrip()[:-1] + line
        elif re.search(r"[^\W\d_]-$", text) and re.match(r"[^\W\d_]", line) and line[0].islower():
            text = text[:-1] + line
        else:
            text = f"{text} {line}"
    return text.replace(_SOFT_HYPHEN, "")


def _pdf_blocks(pages: list[list[_Line]], tables: dict[int, list[tuple[float, float, str]]], sizes: list[tuple[float, float]]) -> list[Block]:
    all_sizes = [round(l.size, 1) for lines in pages for l in lines if l.size > 0]
    body = statistics.median(all_sizes) if all_sizes else 10.0
    heading_sizes = sorted({s for s in all_sizes if s >= body * 1.15}, reverse=True)
    level_of = {s: min(3, i + 1) for i, s in enumerate(heading_sizes)}
    blocks: list[Block] = []
    for index, lines in enumerate(pages):
        page_no = index + 1
        page_tables = sorted(tables.get(index, []), key=lambda t: t[0])
        items: list[tuple[float, str, Any]] = []
        for t0, t1, markdown in page_tables:
            inside = [l for l in lines if t0 - 1 <= (l.top + l.bottom) / 2 <= t1 + 1]
            text = markdown or "\n".join(dehyphenate_join([l.text]) for l in inside)
            if text.strip():
                items.append((t0, "table", text))
        lines = [l for l in lines if not any(t0 - 1 <= (l.top + l.bottom) / 2 <= t1 + 1 for t0, t1, _ in page_tables)]
        items += [(l.top, "line", l) for l in lines]
        items.sort(key=lambda it: it[0])
        gaps = [b.top - a.bottom for a, b in zip(lines, lines[1:]) if b.top > a.bottom]
        typical_gap = statistics.median(gaps) if gaps else 2.0
        current: list[_Line] = []

        def flush() -> None:
            if current:
                text = dehyphenate_join([l.text for l in current])
                if text.strip():
                    blocks.append(Block("paragraph", text, page=page_no))
                current.clear()

        for _top, kind, value in items:
            if kind == "table":
                flush()
                blocks.append(Block("table", value, page=page_no))
                continue
            line: _Line = value
            size = round(line.size, 1)
            is_heading = size in level_of and len(line.text) <= 150 and not line.text.rstrip().endswith((".", ",", ";"))
            if is_heading:
                flush()
                title = re.sub(r"^\d+(?:\.\d+)*\s+", "", line.text).strip() or line.text.strip()
                if blocks and blocks[-1].kind == "heading" and blocks[-1].page == page_no and blocks[-1].level == level_of[size]:
                    blocks[-1].text = f"{blocks[-1].text} {title}"
                else:
                    blocks.append(Block("heading", title, level=level_of[size], page=page_no))
                continue
            if current:
                prev = current[-1]
                gap = line.top - prev.bottom
                size_change = abs(line.size - prev.size) > 0.15 * max(prev.size, 1.0)
                if gap > max(typical_gap * 1.8, 0.6 * prev.height) or size_change:
                    flush()
            current.append(line)
        flush()
    return _join_cross_page_paragraphs(blocks)


def _join_cross_page_paragraphs(blocks: list[Block]) -> list[Block]:
    """A paragraph broken by a page break continues on the next page: glue it
    back (it keeps the page it starts on)."""
    out: list[Block] = []
    for block in blocks:
        prev = out[-1] if out else None
        if (
            prev is not None
            and prev.kind == block.kind == "paragraph"
            and prev.page is not None
            and block.page == prev.page + 1
            and not re.search(r"[.!?:…\"”)]$", prev.text.rstrip())
            and block.text[:1].islower()
        ):
            prev.text = dehyphenate_join([prev.text, block.text])
            continue
        out.append(block)
    return out


def _ocr(image: Any, lang: str, warnings: list[str]) -> str:
    try:
        import pytesseract

        return (pytesseract.image_to_string(image, lang=lang) or "").strip()
    except Exception as exc:  # noqa: BLE001
        warnings.append(f"OCR failed: {exc}")
        return ""


# ------------------------------------------------------------- DOCX / PPTX
def _docx(path: Path) -> list[Block]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    blocks: list[Block] = []
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(element, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            match = re.match(r"(?i)heading\s*(\d)", style)
            if match or style.lower() == "title":
                blocks.append(Block("heading", text, level=int(match.group(1)) if match else 1))
            elif style.lower().startswith("list"):
                blocks.append(Block("list", f"- {text}"))
            else:
                blocks.append(Block("paragraph", text))
        elif tag == "tbl":
            table = Table(element, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            markdown = table_to_markdown(rows)
            if markdown:
                blocks.append(Block("table", markdown))
    return _merge_lists(blocks)


def _pptx(path: Path) -> list[Block]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    blocks: list[Block] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = (title_shape.text or "").strip() if title_shape is not None else ""
        blocks.append(Block("heading", title or f"Slide {number}", level=1, slide=number))
        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                markdown = table_to_markdown(rows)
                if markdown:
                    blocks.append(Block("table", markdown, slide=number))
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                if text:
                    blocks.append(Block("paragraph", text, slide=number))
    return blocks


# ------------------------------------------------------------ MD / TXT
_MD_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$")
_MD_LIST = re.compile(r"^\s*(?:[-*+•]|\d+[.)])\s+")


def parse_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    lines = text.replace("\r\n", "\n").split("\n")
    buffer: list[str] = []
    buffer_kind = "paragraph"

    def flush() -> None:
        nonlocal buffer, buffer_kind
        if buffer:
            joined = "\n".join(buffer) if buffer_kind in {"list", "table", "code"} else " ".join(s.strip() for s in buffer)
            if joined.strip():
                blocks.append(Block(buffer_kind, joined.strip()))
        buffer = []
        buffer_kind = "paragraph"

    in_fence = False
    for line in lines:
        stripped = line.strip()
        if stripped.startswith(("```", "~~~")):
            if in_fence:
                buffer.append(line)
                flush()
                in_fence = False
            else:
                flush()
                in_fence = True
                buffer_kind = "code"
                buffer = [line]
            continue
        if in_fence:
            buffer.append(line)
            continue
        heading = _MD_HEADING.match(line)
        if heading:
            flush()
            blocks.append(Block("heading", heading.group(2).strip(), level=len(heading.group(1))))
            continue
        if not stripped:
            flush()
            continue
        kind = "table" if stripped.startswith("|") else "list" if _MD_LIST.match(line) else "paragraph"
        if buffer and kind != buffer_kind and not (buffer_kind == "list" and kind == "paragraph" and line.startswith(("  ", "\t"))):
            flush()
        if not buffer:
            buffer_kind = kind
        buffer.append(line.rstrip())
    flush()
    return blocks


def parse_plain_text(text: str) -> list[Block]:
    blocks: list[Block] = []
    for para in re.split(r"\n\s*\n", text.replace("\r\n", "\n")):
        lines = [l for l in para.split("\n") if l.strip()]
        if not lines:
            continue
        if all(_MD_LIST.match(l) for l in lines):
            blocks.append(Block("list", "\n".join(l.rstrip() for l in lines)))
        else:
            blocks.append(Block("paragraph", dehyphenate_join(lines)))
    return blocks


def _merge_lists(blocks: list[Block]) -> list[Block]:
    out: list[Block] = []
    for block in blocks:
        if block.kind == "list" and out and out[-1].kind == "list":
            out[-1].text += "\n" + block.text
        else:
            out.append(block)
    return out


def list_kb_files(kb_dir: Path) -> list[Path]:
    return sorted(p for p in kb_dir.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS and not p.name.startswith("."))


def extract_all(
    files: list[Path],
    root: Path,
    extractor: Extractor,
    *,
    max_workers: int = 8,
    on_error: Callable[[Path, Exception], None] | None = None,
) -> list[ExtractedDoc]:
    """Per-file parallel extraction; documents come back in `files` order."""
    from concurrent.futures import ThreadPoolExecutor

    results: list[ExtractedDoc | None] = [None] * len(files)

    def work(i: int) -> None:
        try:
            results[i] = extractor.extract(files[i], root)
        except Exception as exc:  # noqa: BLE001 - one bad file must not sink the KB
            if on_error is not None:
                on_error(files[i], exc)

    workers = max(1, min(max_workers, len(files)))
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="kb-extract") as pool:
        list(pool.map(work, range(len(files))))
    return [doc for doc in results if doc is not None]
