"""Presentation mode (#156): the old CLI's presentation argv, slides as graph
leaves in the book layout, the deck Markdown in the old format, PPTX with
speaker notes, Beamer, narration and TTS (OpenRouter-style endpoint on the
fake LLM; local Coqui through a stand-in `TTS` module), resume rules."""

from __future__ import annotations

import json
import os
import shutil
import sys
import types
import wave
from pathlib import Path

import pytest

from fake_llm import Fault, FakeLLM
from helpers import REPO_ROOT, engine_env, make_work_dir, requires_lualatex, requires_pandoc, run_cli
from responders import node_key_of

STEM = "From_Calculating_Machines_to_Stored_Programs"
CZ_SPEC = REPO_ROOT / "input" / "presentation" / "presentation_input.txt"


def pres_argv(work: Path, *extra: str, input_name: str = "presentation_input.txt") -> list[str]:
    """The presentation argv of the old `main.py --mode presentation`."""
    argv = ["--mode", "presentation", "-i", str(work / input_name), "-o", str(work / "out"), "--use-txt", *extra]
    if (work / "kb").is_dir():
        argv += ["--kb-dir", str(work / "kb")]
    return argv


def pres_run(work: Path, fake: FakeLLM, *extra: str, **env: str):
    return run_cli(pres_argv(work, *extra), fake=fake, env=engine_env(work.parent, **env))


def _leaves(graph: dict) -> list[str]:
    parents = {p for p, _c in graph["edges"]}
    return [k for k in graph["nodes"] if k != "book" and k not in parents]


def _tts_calls(fake: FakeLLM) -> list:
    return [c for c in fake.calls if c.path.endswith("/audio/speech")]


class FakeCoqui:
    """Stands in for `TTS.api.TTS`: writes a short 16 kHz mono WAV per call."""

    instances: list["FakeCoqui"] = []

    def __init__(self, model_name: str, progress_bar: bool = False, gpu: bool = False) -> None:
        self.model_name = model_name
        self.calls: list[dict] = []
        self.speakers = None
        FakeCoqui.instances.append(self)

    def tts_to_file(self, text: str, file_path: str, **kwargs) -> None:
        self.calls.append({"text": text, **kwargs})
        with wave.open(file_path, "wb") as out:
            out.setnchannels(1)
            out.setsampwidth(2)
            out.setframerate(16000)
            out.writeframes(b"\x00\x00" * int(16000 * 0.05 * len(text.split())))


@pytest.fixture
def fake_coqui(monkeypatch):
    FakeCoqui.instances = []
    package = types.ModuleType("TTS")
    api = types.ModuleType("TTS.api")
    api.TTS = FakeCoqui
    package.api = api
    monkeypatch.setitem(sys.modules, "TTS", package)
    monkeypatch.setitem(sys.modules, "TTS.api", api)
    return FakeCoqui


def test_presentation_argv_is_accepted_and_unknown_flags_exit_2() -> None:
    from engine.cli import build_parser

    args = build_parser().parse_args([
        "--mode", "presentation", "--presentation-input", "p.txt", "-o", "out", "--presentation-tex", "--presentation-pptx",
        "--presentation-narration", "--presentation-narration-model", "m", "--presentation-tts", "--presentation-tts-mode", "local",
        "--presentation-tts-model", "t", "--presentation-exclude-slides", "2,4-5", "--disable-general-knowledge-citation",
        "--kb-dir", "kb", "--use-json", "-j", "presentation_structure.json", "--resume", "--no-pdf",
    ])
    assert args.presentation_tts_mode == "local" and args.presentation_exclude_slides == "2,4-5"
    for bad in (["--presentation-tts-mode", "cloud"], ["--presentation-video"]):
        with pytest.raises(SystemExit) as exc:
            build_parser().parse_args(["--mode", "presentation", *bad])
        assert exc.value.code == 2


def test_full_presentation_run(tmp_path: Path) -> None:
    work = make_work_dir(tmp_path, bench="en_presentation")
    fake = FakeLLM()
    run = pres_run(work, fake, "--presentation-pptx")
    assert run.exit_code == 0, run.text
    out = run.out_dir
    graph = run.graph()
    assert graph["graph"]["doc_type"] == "presentation" and graph["graph"]["theme"] == "Madrid"
    leaves = _leaves(graph)
    assert len(leaves) == 6 and "3-1" in leaves and "3-2" in leaves  # the two-slide block was split
    assert sorted(p.stem for p in (out / "sections").glob("*.md")) == sorted(leaves)
    assert all(Path(graph["nodes"][k]["content_file_path"]) == out / "sections" / f"{k}.md" for k in leaves)
    structure = json.loads((out / "presentation_structure.json").read_text(encoding="utf-8"))
    assert structure["duration_minutes"] == 10 and [s["n_pages"] for s in structure["slides"]] == [1, 1, 2, 1, 1]
    meta = run.run_meta()
    assert meta["mode"] == "presentation" and meta["run_kind"] == "full" and meta["error"] is None
    # The deck Markdown keeps the old engine's format.
    deck = (out / f"{STEM}.md").read_text(encoding="utf-8")
    assert deck.startswith("theme: Madrid\npaginate: false\noutline: true\n\n---\n<!-- _class: title -->\n# From Calculating Machines")
    frames = deck.split("\n---\n")[1:]
    assert len(frames) == 1 + 6 + 1  # title, slides, sources
    assert frames[-1].startswith("# Sources") and "- [1] " in frames[-1]
    assert "[kb_" not in deck and "[1]" in frames[1]
    for key in leaves:
        body = (out / "sections" / f"{key}.md").read_text(encoding="utf-8")
        assert "[kb_" in body  # slide files keep cite keys (what the API imports)
        assert sum(1 for line in body.splitlines() if line.strip()) <= 10
    # Writers ran in parallel over the slides, each seeing its neighbours' plan.
    [call] = [c for c in fake.chat_calls("SlideDraft") if node_key_of(c.prompt) == "3-2"]
    assert "Previous slide (Programs on cards, part 1)" in call.user and "Must include" in call.user
    # PPTX: one slide per frame, titles in order.
    from pptx import Presentation

    pptx = Presentation(str(out / f"{STEM}.pptx"))
    titles = [s.shapes.title.text for s in pptx.slides]
    assert len(titles) == 8 and titles[0].startswith("From Calculating") and titles[-1] == "Sources"
    assert not fake.chat_calls("SlideNarration") and not _tts_calls(fake)


def test_slide_bodies_keep_bold_label_lines_separate(tmp_path: Path) -> None:
    """The presentation pipeline cleans slide drafts with `lead_ins=False`: a bold
    label on its own line is not merged into the next line (the document rule)."""
    work = make_work_dir(tmp_path, bench="en_presentation")
    reply = {"body_markdown": "**Takeaway**\nShort statement.", "summary": "s", "citations_used": []}
    fake = FakeLLM(overrides={"SlideDraft": lambda call: reply})
    run = pres_run(work, fake)
    assert run.exit_code == 0, run.text
    bodies = [p.read_text(encoding="utf-8") for p in (run.out_dir / "sections").glob("*.md")]
    assert bodies and all(b.strip() == "**Takeaway**\nShort statement." for b in bodies)


def test_explicit_czech_slides_are_used_verbatim(tmp_path: Path) -> None:
    work = tmp_path / "cz"
    work.mkdir()
    shutil.copyfile(CZ_SPEC, work / "presentation_input.txt")
    fake = FakeLLM()
    run = pres_run(work, fake)
    assert run.exit_code == 0, run.text
    assert not fake.chat_calls("PresentationOutline") and not fake.chat_calls("SubdivisionPlan")
    graph = run.graph()
    assert graph["graph"]["language"] == "cs" and graph["nodes"]["book"]["title"] == "Kurz pedagogických dovedností"
    assert [graph["nodes"][k]["title"] for k in _leaves(graph)][:3] == ["Úvod", "Kontext a obsah kurzu", "Téma 1"]
    [call] = [c for c in fake.chat_calls("SlideDraft") if node_key_of(c.prompt) == "3"]
    assert "- 1. Vzdělávací cíle" in call.user and "Czech" in call.system  # the speaker's draft and the output language


def test_narration_and_openrouter_tts(tmp_path: Path) -> None:
    work = make_work_dir(tmp_path, bench="en_presentation")
    fake = FakeLLM()
    run = pres_run(work, fake, "--presentation-tts", "--presentation-exclude-slides", "2,4-5")
    assert run.exit_code == 0, run.text
    out = run.out_dir
    narration = json.loads((out / f"{STEM}_narration.json").read_text(encoding="utf-8"))["slides"]
    assert [item["index"] for item in narration] == [1, 2, 3, 4, 5, 6, 7]  # the title slide is slide 1
    assert narration[0]["title"].startswith("From Calculating") and narration[1]["title"] == "Motivation"
    script = (out / f"{STEM}_narration.md").read_text(encoding="utf-8")
    assert script.startswith("# Narration Script\n\n## From Calculating") and "## Motivation\n\nFake narration for slide 2" in script
    # Audio: every slide except the excluded 2, 4 and 5.
    assert sorted(p.name for p in (out / "audio").iterdir()) == ["slide_01.mp3", "slide_03.mp3", "slide_06.mp3", "slide_07.mp3"]
    calls = _tts_calls(fake)
    assert len(calls) == 4 and {c.body["model"] for c in calls} == {"openai/gpt-4o-mini-tts-2025-12-15"}
    assert {c.body["voice"] for c in calls} == {"alloy"} and {c.body["response_format"] for c in calls} == {"mp3"}
    from engine.media.audio import mp3_duration

    clips = [mp3_duration((out / "audio" / name).read_bytes()) for name in ("slide_01.mp3", "slide_03.mp3", "slide_06.mp3", "slide_07.mp3")]
    combined = mp3_duration((out / f"{STEM}_audio.mp3").read_bytes())
    assert abs(combined - sum(clips)) < 0.05
    subtitles = (out / f"{STEM}.srt").read_text(encoding="utf-8")
    assert subtitles.startswith("1\n00:00:00,000 --> 00:00:") and subtitles.count(" --> ") == 4
    assert "slide 2," not in subtitles and "slide 3," in subtitles
    tts_usage = [u for u in run.usage_lines() if u["kind"] == "tts"]
    assert len(tts_usage) == 4 and {u["label"] for u in tts_usage} == {"presentation.tts"}
    assert all(u["cost_usd"] is None for u in tts_usage)  # no token-based pricing for audio
    # Narration is the speaker notes of the PPTX when both are requested.
    fake2 = FakeLLM()
    again = pres_run(work, fake2, "--resume", "--presentation-tts", "--presentation-exclude-slides", "2,4-5", "--presentation-pptx")
    assert again.exit_code == 0 and fake2.calls == []  # narration and audio are cached
    assert again.run_meta()["run_kind"] == "export"
    from pptx import Presentation

    notes = [s.notes_slide.notes_text_frame.text for s in Presentation(str(out / f"{STEM}.pptx")).slides]
    assert notes[1].startswith("Fake narration for slide 2")


def test_local_tts_with_coqui(tmp_path: Path, fake_coqui) -> None:
    work = make_work_dir(tmp_path, bench="en_presentation")
    run = pres_run(work, FakeLLM(), "--presentation-tts", "--presentation-tts-mode", "local")
    assert run.exit_code == 0, run.text
    out = run.out_dir
    [model] = fake_coqui.instances  # loaded once, reused for every slide
    assert model.model_name == "tts_models/en/ljspeech/vits" and len(model.calls) == 7
    assert sorted(p.suffix for p in (out / "audio").iterdir()) == [".wav"] * 7
    with wave.open(str(out / f"{STEM}_audio.wav"), "rb") as handle:
        assert handle.getframerate() == 16000 and handle.getnframes() > 0
    assert (out / f"{STEM}.srt").read_text(encoding="utf-8").count(" --> ") == 7
    # A speaker reference switches to XTTS v2 with the output language.
    work2 = make_work_dir(tmp_path, bench="en_presentation", name="xtts")
    reference = tmp_path / "voice.wav"
    shutil.copyfile(out / "audio" / "slide_01.wav", reference)
    run2 = pres_run(work2, FakeLLM(), "--presentation-tts", "--presentation-tts-mode", "local", AUTOGENBOOK_TTS_SPEAKER_WAV=str(reference))
    assert run2.exit_code == 0, run2.text
    xtts = fake_coqui.instances[-1]
    assert xtts.model_name.endswith("xtts_v2") and xtts.calls[0]["language"] == "en" and xtts.calls[0]["speaker_wav"] == [str(reference)]


def test_local_tts_without_coqui_fails_cleanly(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setitem(sys.modules, "TTS", None)  # import fails as if the extra were not installed
    work = make_work_dir(tmp_path, bench="en_presentation")
    run = pres_run(work, FakeLLM(), "--presentation-tts", "--presentation-tts-mode", "local")
    assert run.exit_code == 1
    assert "pip install '.[tts-local]'" in run.run_meta()["error"]
    assert (run.out_dir / f"{STEM}.md").exists() and (run.out_dir / f"{STEM}_narration.json").exists()


def test_tts_failure_is_reported_and_retried_on_resume(tmp_path: Path) -> None:
    work = make_work_dir(tmp_path, bench="en_presentation")
    failing = FakeLLM(faults=[Fault(status=400, match=lambda c: c.path.endswith("/audio/speech") and "slide 3," in str(c.body.get("input")), times=None)])
    first = pres_run(work, failing, "--presentation-tts")
    assert first.exit_code == 1 and "400" in first.run_meta()["error"]
    assert not (first.out_dir / "audio" / "slide_03.mp3").exists()
    fake = FakeLLM()
    retry = pres_run(work, fake, "--resume", "--presentation-tts")
    assert retry.exit_code == 0, retry.text
    assert retry.run_meta()["run_kind"] == "resume"
    assert not fake.chat_calls()  # slides and narration were kept
    assert [c.body["input"].split(",")[0] for c in _tts_calls(fake)] == ["Fake narration for slide 3"]


def test_regenerate_one_slide_redoes_only_its_media(tmp_path: Path) -> None:
    work = make_work_dir(tmp_path, bench="en_presentation")
    assert pres_run(work, FakeLLM(), "--presentation-tts").exit_code == 0
    out = work / "out"
    (out / "sections" / "3-1.md").unlink()
    graph = json.loads((out / "structure_graph.json").read_text(encoding="utf-8"))
    graph["nodes"]["3-1"]["content_file_path"] = ""
    (out / "structure_graph.json").write_text(json.dumps(graph), encoding="utf-8")
    new_text = {"body_markdown": "- A rewritten point\n- Another one", "summary": "Rewritten.", "citations_used": []}
    fake = FakeLLM(overrides={
        "SlideDraft": lambda call: new_text,
        "SlideNarration": lambda call: {"narration": "New narration: " + call.user.split("\n- ", 1)[-1].split("\n")[0]},
    })
    run = pres_run(work, fake, "--resume", "--presentation-tts")
    assert run.exit_code == 0, run.text
    assert [node_key_of(c.prompt) for c in fake.chat_calls("SlideDraft")] == ["3-1"]
    # Only the changed slide gets new narration and audio (keyed by its text).
    [narration] = fake.chat_calls("SlideNarration")
    assert "A rewritten point" in narration.user and len(_tts_calls(fake)) == 1
    # Narration requested later for an existing deck: only the narration is made.
    work2 = make_work_dir(tmp_path, bench="en_presentation", name="later")
    assert pres_run(work2, FakeLLM()).exit_code == 0
    fake2 = FakeLLM()
    later = pres_run(work2, fake2, "--resume", "--presentation-narration")
    assert later.exit_code == 0 and later.run_meta()["run_kind"] == "resume"
    assert fake2.schema_counts() == {"SlideNarration": 7}


def test_old_engine_structure_json_loads(tmp_path: Path) -> None:
    """-j presentation_structure.json --use-json with the old engine's file."""
    work = tmp_path / "old"
    (work / "out").mkdir(parents=True)
    shutil.copyfile(CZ_SPEC, work / "presentation_input.txt")
    shutil.copyfile(REPO_ROOT / "output" / "presentation" / "presentation_structure.json", work / "out" / "presentation_structure.json")
    fake = FakeLLM()
    run = run_cli(["--mode", "presentation", "-i", str(work / "presentation_input.txt"), "-o", str(work / "out"),
                   "-j", "presentation_structure.json", "--use-json"], fake=fake, env=engine_env(tmp_path))
    assert run.exit_code == 0, run.text
    assert not fake.chat_calls("PresentationOutline")
    graph = run.graph()
    assert graph["nodes"]["book"]["title"] == "Kurz pedagogických dovedností" and len(_leaves(graph)) == 10


@requires_pandoc
def test_beamer_tex(tmp_path: Path) -> None:
    work = make_work_dir(tmp_path, bench="en_presentation")
    run = pres_run(work, FakeLLM(), "--presentation-tex", "--no-pdf")
    assert run.exit_code == 0, run.text
    tex = (run.out_dir / f"{STEM}.tex").read_text(encoding="utf-8")
    assert "\\documentclass" in tex and "beamer" in tex and "\\usetheme[]{Madrid}" in tex
    assert tex.count("\\begin{frame}") >= 7 and "\\tableofcontents" in tex
    assert not (run.out_dir / f"{STEM}.pdf").exists()


@requires_lualatex
def test_beamer_pdf(tmp_path: Path) -> None:
    work = make_work_dir(tmp_path, bench="en_presentation")
    run = pres_run(work, FakeLLM(), "--presentation-tex")
    assert run.exit_code == 0, run.text
    pdf = run.out_dir / f"{STEM}.pdf"
    assert pdf.exists() and pdf.stat().st_size > 5000
    import pypdfium2 as pdfium

    assert len(pdfium.PdfDocument(str(pdf))) >= 9  # title, outline, 6 slides, sources


@pytest.mark.skipif(not os.environ.get("ENGINE_TEST_COQUI"), reason="set ENGINE_TEST_COQUI=1 with a downloaded Coqui model")
def test_real_coqui_synthesis(tmp_path: Path) -> None:
    """Local synthesis with the real Coqui package. Opt-in: loading a model
    may download it, and tests never touch the network by default; CI does
    not install the tts-local extra."""
    pytest.importorskip("TTS")
    import asyncio

    from engine.media.tts import LocalCoquiTTS

    provider = LocalCoquiTTS(language="en")
    try:
        clip = asyncio.run(provider.synthesize("A short test sentence.", tmp_path / "clip.wav"))
    except Exception as exc:  # noqa: BLE001 - no model download possible offline
        pytest.skip(f"Coqui model unavailable: {exc}")
    assert clip.seconds > 0
